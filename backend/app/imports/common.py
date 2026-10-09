"""Gemeinsame Hilfen für CSV-Importe: Einlesen, Ergebnisobjekt, Ablage und Protokoll."""
import csv
import io
import re
import uuid
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from app.config import settings


class ImportAborted(Exception):
    """Der gesamte Import kann nicht durchgeführt werden (nichts wird gespeichert)."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


class RowError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


@dataclass
class ImportResult:
    total: int = 0
    imported: int = 0
    updated: int = 0
    errors: list[dict] = field(default_factory=list)
    created_kreise: list[dict] = field(default_factory=list)


def decode(raw: bytes) -> str:
    """UTF-8 (mit oder ohne BOM), sonst Windows-1250 (polnisches Excel)."""
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1250")


def parse_csv(text: str, skip_header: bool = True) -> list[tuple[int, list[str]]]:
    """Liefert (Zeilennummer, Zellen), ohne Leerzeilen und optional ohne Überschriftszeile."""
    first_line = text.lstrip("\ufeff").split("\n", 1)[0]
    counts = {d: first_line.count(d) for d in (";", ",", "\t")}
    delimiter = max(counts, key=counts.get) if any(counts.values()) else ";"

    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    rows = []
    header_seen = not skip_header
    for cells in reader:
        if not any(c.strip() for c in cells):
            continue
        if not header_seen:
            header_seen = True
            continue
        rows.append((reader.line_num, cells))
    return rows


def read_csv_rows(raw: bytes, filename: str | None = None) -> list[tuple[int, list[str]]]:
    """CSV ohne Überschriftszeile."""
    try:
        return parse_csv(decode(raw))
    except UnicodeDecodeError as exc:
        raise ImportAborted("unreadable", "Datei kann nicht gelesen werden") from exc


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value).strip()


def read_table(raw: bytes, filename: str | None = None) -> list[tuple[int, list[str]]]:
    """Alle nicht leeren Zeilen einer Excel-Datei (erstes Blatt) oder CSV-/Textdatei, mit Zeilennummer."""
    if (filename or "").lower().endswith(".xlsx") or raw[:2] == b"PK":
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            rows = []
            for number, row in enumerate(workbook.worksheets[0].iter_rows(values_only=True), start=1):
                cells = [cell_text(c) for c in row]
                if any(cells):
                    rows.append((number, cells))
            workbook.close()
            return rows
        except Exception as exc:
            raise ImportAborted("unreadable", "Excel-Datei kann nicht gelesen werden") from exc
    try:
        return parse_csv(decode(raw), skip_header=False)
    except UnicodeDecodeError as exc:
        raise ImportAborted("unreadable", "Datei kann nicht gelesen werden") from exc


def protocol_text(key: str, filename: str, result: ImportResult) -> str:
    lines = [
        f"{key} - {filename}",
        f"Zeilen: {result.total}, importiert: {result.imported}, Fehler: {len(result.errors)}",
    ]
    if result.updated:
        lines.append(f"Davon vorhandene Einträge aktualisiert: {result.updated}")
    if result.created_kreise:
        lines.append("Neu angelegte Kreise (polnisch): " + ", ".join(k["name"] for k in result.created_kreise))
    for error in result.errors:
        lines.append(f"Zeile {error['line']}: {error['message']} | {';'.join(error['values'])}")
    return "\n".join(lines) + "\n"


def archive(raw: bytes, key: str, filename: str | None, protocol: str) -> str:
    """Speichert die abgearbeitete Datei samt Protokoll im konfigurierten Ordner."""
    directory = Path(settings.import_archive_dir)
    directory.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^A-Za-z0-9._-]", "_", Path(filename or "upload.csv").name)[:100]
    name = f"{datetime.now().strftime('%Y%m%d-%H%M%S')}_{key}_{uuid.uuid4().hex[:6]}_{safe}"
    (directory / name).write_bytes(raw)
    (directory / f"{name}.log").write_text(protocol, encoding="utf-8")
    return name
